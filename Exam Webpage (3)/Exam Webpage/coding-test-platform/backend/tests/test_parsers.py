"""Question-document and student-list parsing. Run: python -m pytest tests/test_parsers.py -q"""
import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.reios.parsers import ParseError, parse_question_file, student_rows


def docx_bytes(build) -> bytes:
    d = Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def answers(qs):
    return ["".join("ABCDEFGH"[i] for i in q["correct_options"]) for q in qs]


def test_typed_word_document():
    lines = ["Mock set", "Section: Quantitative Aptitude",
             "1. What is 20% of 150?", "A) 20", "B) 25", "C) 30", "D) 35", "Answer: C",
             "Explanation: 20/100 x 150", "2. Speed of 180 km in 3 h in m/s",
             "(a) 15  (b) 16.67  (c) 18  (d) 20", "Ans: (b)",
             "Section: General Knowledge", "Q3. Red planet?", "A. Venus", "B. Mars *", "C. Jupiter",
             "4. Which are prime?", "A) 2", "B) 9", "C) 11", "Correct answer: A and C"]
    data = docx_bytes(lambda d: [d.add_paragraph(t) for t in lines])
    qs, warnings = parse_question_file(data, "set.docx")
    assert not warnings and len(qs) == 4
    assert answers(qs) == ["C", "B", "B", "AC"]
    assert [q["section"] for q in qs] == ["Quantitative Aptitude", "Quantitative Aptitude",
                                          "General Knowledge", "General Knowledge"]
    assert qs[0]["explanation"] == "20/100 x 150"
    assert qs[1]["options"] == ["15", "16.67", "18", "20"]
    assert qs[2]["options"][1] == "Mars"
    assert qs[3]["is_multi"] and all(not q["problems"] for q in qs)


def test_auto_numbered_word_with_answer_key():
    def build(d):
        for q, opts in [("Unit of force", ["Joule", "Newton", "Watt"]), ("Vector quantity", ["Mass", "Velocity"])]:
            d.add_paragraph(q, style="List Number")
            for o in opts:
                d.add_paragraph(o, style="List Number 2")
        d.add_paragraph("Answer Key")
        d.add_paragraph("1. B  2. B")
    qs, _ = parse_question_file(docx_bytes(build), "physics.docx", default_section="Physics")
    assert [q["question_text"] for q in qs] == ["Unit of force", "Vector quantity"]
    assert [len(q["options"]) for q in qs] == [3, 2]
    assert answers(qs) == ["B", "B"] and qs[0]["section"] == "Physics"


def test_missing_answer_is_flagged_not_dropped():
    qs, _ = parse_question_file(b"1. Q one\nA) x\nB) y\n2. Q two\nA) p\nB) q\nAnswer: A\n", "set.txt")
    assert len(qs) == 2
    assert any("Correct answer" in p for p in qs[0]["problems"]) and not qs[1]["problems"]


def test_excel_questions_and_bad_files():
    wb = Workbook()
    wb.active.append(["section", "question", "option_a", "option_b", "option_c", "correct", "marks"])
    wb.active.append(["Reasoning", "Odd one out", "Apple", "Carrot", "Mango", "B", 2])
    buf = io.BytesIO()
    wb.save(buf)
    qs, _ = parse_question_file(buf.getvalue(), "set.xlsx")
    assert answers(qs) == ["B"] and qs[0]["marks"] == 2.0
    with pytest.raises(ParseError):
        parse_question_file(b"not a zip", "set.docx")
    with pytest.raises(ParseError):
        parse_question_file(b"x", "old.doc")


def test_student_list_from_word_table_with_college_headers():
    def build(d):
        t = d.add_table(rows=1, cols=3)
        for c, h in zip(t.rows[0].cells, ["Roll Number", "Student Name", "Password"]):
            c.text = h
        for r in (["22CS1", "Ravi", "Pass@123"], ["22CS2", "Sneha", ""]):
            for c, v in zip(t.add_row().cells, r):
                c.text = v
    rows = student_rows(docx_bytes(build), "students.docx")
    assert rows == [{"roll_no": "22CS1", "name": "Ravi", "password": "Pass@123"},
                    {"roll_no": "22CS2", "name": "Sneha"}]
    with pytest.raises(ParseError):
        student_rows(b"name,email\nA,a@x\n", "s.csv")


def test_one_file_holding_several_sets():
    text = ("Section: General Knowledge – Set 1\n"
            "1. First?\nA) a\nB) b\nAnswer: A\n"
            "2. Founder of INC?\nA) Naoroji\nB) W. C. Bonnerjee\nC) Hume\nAnswer: C\n"
            "\u00a0\n"
            "Section: General Knowledge – Set 2\n"
            "1. Kalam?\nA) x\nB) A. P. J. Abdul Kalam\nAnswer: B\n"
            "Set 3\n"
            "1. Sodium?\nA) Sd  B) Na  C) S\nAnswer: B\n")
    qs, warnings = parse_question_file(text.encode(), "all_sets.txt")
    assert not warnings
    assert [(q["set"], q["number"]) for q in qs] == [("Set 1", 1), ("Set 1", 2), ("Set 2", 1), ("Set 3", 1)]
    assert {q["section"] for q in qs} == {"General Knowledge"}
    assert qs[1]["options"] == ["Naoroji", "W. C. Bonnerjee", "Hume"] and answers(qs) == ["A", "C", "B", "B"]
    assert qs[2]["options"][1] == "A. P. J. Abdul Kalam" and qs[3]["options"] == ["Sd", "Na", "S"]


def test_answer_key_numbers_restart_per_set():
    text = "Set 1\n1. Q\nA) x\nB) y\nSet 2\n1. Q\nA) x\nB) y\nAnswer Key\nSet 1\n1. B\nSet 2\n1. A\n"
    qs, _ = parse_question_file(text.encode(), "k.txt")
    assert answers(qs) == ["B", "A"]
